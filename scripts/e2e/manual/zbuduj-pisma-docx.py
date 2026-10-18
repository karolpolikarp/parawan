# Buduje decyzja.docx i protokol-kontroli.docx z fikcyjnych pism .txt w apps/web/src/fixtures/urzad.
# python-docx daje pakiet jak z Worda (style, motyw, ustawienia, customXml, miniatura), a nie
# minimalny ZIP z samym document.xml, więc testy importu i zapisu DOCX widzą prawdziwe części.
# Użycie (z katalogu repo): python scripts/e2e/manual/zbuduj-pisma-docx.py [katalog-wyjściowy]
# Pisma .txt są zawsze czytane z apps/web/src/fixtures/urzad, a gotowe .docx trafiają do katalogu
# z argumentu (domyślnie ten sam katalog w repo, czyli nadpisanie plików). Dawniej skrypt szukał
# .txt w katalogu z argumentu i z pustym katalogiem kończył się błędem braku pliku.
# Wymaga: pip install python-docx (sprawdzone na 1.2.0).
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[3]
SRC = ROOT / "apps" / "web" / "src" / "fixtures" / "urzad"
F = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else SRC
F.mkdir(parents=True, exist_ok=True)


def nowy_dokument():
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    return d


# Protokół: lista obecnych jako prawdziwa tabela Worda (wiersze z tabulatorami w .txt).
src = (SRC / "protokol-kontroli.txt").read_text(encoding="utf-8").split("\n")
d = nowy_dokument()
i = 0
while i < len(src):
    line = src[i]
    if line.startswith("Lp.\t"):
        rows = []
        while i < len(src) and "\t" in src[i]:
            rows.append(src[i].split("\t"))
            i += 1
        t = d.add_table(rows=len(rows), cols=3)
        t.style = "Table Grid"
        for r, row in enumerate(rows):
            for c, val in enumerate(row[:3]):
                t.cell(r, c).text = val
        continue
    p = d.add_paragraph(line)
    if line.startswith("PROTOKÓŁ"):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].bold = True
    i += 1
d.save(F / "protokol-kontroli.docx")

# Decyzja: nagłówek urzędu w tabeli dwukolumnowej, jak w typowym szablonie pisma.
src = (SRC / "decyzja.txt").read_text(encoding="utf-8").split("\n")
d = nowy_dokument()
t = d.add_table(rows=1, cols=2)
t.cell(0, 0).text = "\n".join(src[0:4])
t.cell(0, 1).text = src[5]
for line in src[6:]:
    p = d.add_paragraph(line)
    if line.startswith("/-/") or line.startswith("Wójt Gminy Brzozowa Wola"):
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
d.save(F / "decyzja.docx")
print("ok:", F)
